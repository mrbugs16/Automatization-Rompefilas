"""Lee la matriz de pruebas (data/RompeFilas.xlsx) y escribe una copia con los resultados obtenidos."""
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from config import settings
from core.xlsx import guardar_con_graficas

COLUMNAS = {
    "Producto": "producto",
    "Módulo": "modulo",
    "Escenario": "escenario",
    "ID Caso": "id",
    "Nombre del Caso de prueba": "caso",
    "Descripción del Caso de Prueba": "descripcion",
    "Pre condiciónes": "precondiciones",
    "Pasos de Ejecución": "pasos",
    "Resultado Esperado": "esperado",
    "Prioridad": "prioridad",
    "Estatus Automatización": "automatizacion",
    "Tipo de Ejecución": "ejecucion",
    "Tipo de Prueba": "tipo",
    "Observaciones": "observaciones",
}

# Estados del reporte -> valores permitidos en la columna 'Estatus' (validación de datos del Excel)
ESTADO_EXCEL = {
    "Aprobado": "Pasó",
    "Fallido": "Falló",
    "Error": "Bloqueado",
    "Pendiente": "Bloqueado",
    "Manual": "Sin Ejecutar",
}

COLORES_ESTADO = {
    "Pasó": "C6EFCE",
    "Falló": "FFC7CE",
    "Bloqueado": "FFEB9C",
}


def _separar_datos(precondiciones):
    """La celda trae 'precondición ... Datos de prueba: ...': se separa para el reporte."""
    texto = str(precondiciones or "")
    if "Datos de prueba:" in texto:
        antes, despues = texto.split("Datos de prueba:", 1)
        return antes.strip(), despues.strip()
    return texto.strip(), ""


def cargar_matriz(ruta=settings.MATRIZ_XLSX):
    """Regresa {ID: {campo: valor}} en el orden de la matriz."""
    ws = load_workbook(ruta, read_only=True, data_only=True)[settings.HOJA_MATRIZ]
    filas = ws.iter_rows(min_row=settings.FILA_ENCABEZADOS, values_only=True)
    encabezados = [COLUMNAS.get(str(h).strip()) if h else None for h in next(filas)]
    casos = {}
    for fila in filas:
        registro = {k: (v if v is not None else "") for k, v in zip(encabezados, fila) if k}
        if str(registro.get("id", "")).startswith("CP-") and registro.get("caso"):
            registro["precondiciones"], registro["datos"] = _separar_datos(registro.get("precondiciones"))
            casos[registro["id"]] = registro
    return casos


def _observacion(r):
    """Resultado obtenido + referencia y monto usados (si el caso llegó a consultar una referencia)."""
    datos = [f"Ref: {r['referencia']}" if r.get("referencia") else "",
             f"Monto: ${r['monto']:,.2f}" if r.get("monto") else "",
             "6 MSI" if r.get("tipo_pago") == "6 MSI" else ""]
    datos = " · ".join(d for d in datos if d)
    return f"{r['obtenido']} [{datos}]" if datos else r["obtenido"]


def escribir_resultados(resultados, destino, origen=settings.MATRIZ_XLSX):
    """Copia la matriz a 'destino' llenando Estatus, Tipo de Ejecución, Fecha, Evidencia y Observaciones."""
    wb = load_workbook(origen)
    ws = wb[settings.HOJA_MATRIZ]
    fila_enc = settings.FILA_ENCABEZADOS
    col = {str(c.value).strip(): c.column for c in ws[fila_enc] if c.value}
    hoy = datetime.now()
    for fila in range(fila_enc + 1, ws.max_row + 1):
        caso_id = ws.cell(fila, col["ID Caso"]).value
        r = resultados.get(caso_id)
        if not r:
            continue
        estado = ESTADO_EXCEL.get(r["estado"], r["estado"])
        celda = ws.cell(fila, col["Estatus"])
        celda.value = estado
        if estado in COLORES_ESTADO:
            color = COLORES_ESTADO[estado]
            celda.fill = PatternFill("solid", start_color=color, end_color=color)

        if r["estado"] == "Manual":
            ws.cell(fila, col["Observaciones"]).value = f"[Prueba manual] {_observacion(r)}"
            continue

        ws.cell(fila, col["Estatus Automatización"]).value = "Automatizada"
        ws.cell(fila, col["Tipo de Ejecución"]).value = "Automatizada"
        fecha = ws.cell(fila, col["Fecha de Ejecución"])
        fecha.value = hoy
        fecha.number_format = "dd/mm/yyyy"
        evidencia = ws.cell(fila, col["Evidencia"])
        evidencia.value = f"reporte.html#{caso_id}"
        evidencia.hyperlink = f"reporte.html#{caso_id}"
        ws.cell(fila, col["Observaciones"]).value = _observacion(r)
    guardar_con_graficas(wb, destino, origen)
